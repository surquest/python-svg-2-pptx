import io
import logging
from typing import Dict
from lxml import etree as lxml_etree

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from ..models import (
    IRSlide, IRNode, IRInfoBox, IRText, IRRectangle, IRIcon, IRLine, IRGroup,
    IREllipse, IRPolygon, IRImage,
    IRConnector, Point, Color, GradientFill, ArrowType, ConnectorType, TextBlock
)
from .image_patch import _pptx_image

logger = logging.getLogger(__name__)

class PPTXBackend:
    """Backend: IR -> python-pptx Presentation."""

    def __init__(self, ir_slides):
        if not isinstance(ir_slides, list):
            ir_slides = [ir_slides]
        self.ir_slides = ir_slides
        self.prs = Presentation()
        if self.ir_slides:
            self.prs.slide_width = self.ir_slides[0].width_emu
            self.prs.slide_height = self.ir_slides[0].height_emu

    def render(self) -> Presentation:
        blank_layout = self.prs.slide_layouts[6]
        for ir_slide in self.ir_slides:
            self.current_ir_slide = ir_slide
            self.slide = self.prs.slides.add_slide(blank_layout)
            # Layout registry maps shape_ids/refs to pptx shape objects
            self._registry: Dict[str, object] = {}
            self._connectable_shapes = []

            # Pass 1a: Render shapes, infoBoxes, icons (non-text). Build registry.
            for node in ir_slide.nodes:
                if not isinstance(node, (IRConnector, IRText)):
                    self._dispatch_shape(node)

            # Pass 1b: Render texts (z-order top)
            for node in ir_slide.nodes:
                if isinstance(node, IRText):
                    self._dispatch_shape(node)

            # Pass 2: Render connectors (now that all shapes exist)
            for node in ir_slide.nodes:
                if isinstance(node, IRConnector):
                    self._render_connector(node)
        return self.prs

    # ---------- Pass 1 dispatch ----------

    def _dispatch_shape(self, node: IRNode):
        if isinstance(node, IRInfoBox):
            return self._render_infobox(node)
        elif isinstance(node, IRText):
            return self._render_text_freeform(node)
        elif isinstance(node, IRRectangle):
            return self._render_rectangle(node)
        elif isinstance(node, IREllipse):
            return self._render_ellipse(node)
        elif isinstance(node, IRPolygon):
            return self._render_polygon(node)
        elif isinstance(node, IRIcon):
            return self._render_icon(node)
        elif isinstance(node, IRLine):
            return self._render_decorative_line(node)
        elif isinstance(node, IRGroup):
            return self._render_group(node)
        elif isinstance(node, IRImage):
            return self._render_image(node)
        return None

    # ---------- Renderers ----------

    def _embed_text_in_shape(self, shape, ir_text: IRText):
        tf = shape.text_frame
        tf.clear()
        tf.word_wrap = True
        
        rect_x = Emu(int(shape.left))
        rect_y = Emu(int(shape.top))
        rect_w = Emu(int(shape.width))
        text_x = ir_text.block.anchor_x
        text_y = ir_text.block.anchor_y
        
        if ir_text.block.alignment == PP_ALIGN.LEFT:
            tf.margin_left = max(Emu(0), text_x - rect_x)
            tf.margin_right = Emu(0)
        elif ir_text.block.alignment == PP_ALIGN.RIGHT:
            tf.margin_left = Emu(0)
            tf.margin_right = max(Emu(0), (rect_x + rect_w) - text_x)
        else:
            diff = text_x - (rect_x + rect_w / 2)
            if diff > 0:
                tf.margin_left = Emu(int(diff * 2))
                tf.margin_right = Emu(0)
            else:
                tf.margin_left = Emu(0)
                tf.margin_right = Emu(int(-diff * 2))
        
        first_font_size = ir_text.block.runs[0].font.size_pt if ir_text.block.runs else 12
        font_size_emu = Emu(int(first_font_size * 12700))
        tf.margin_top = max(Emu(0), (text_y - rect_y) - font_size_emu)
        tf.margin_bottom = Emu(0)

        tf.vertical_anchor = MSO_ANCHOR.TOP
        self._populate_textframe(tf, ir_text.block)

    def _resolve_group_texts(self, ppt_shapes, text_children):
        if text_children and ppt_shapes:
            auto_shapes = [s for s in ppt_shapes if getattr(s, "has_text_frame", False)]
            if auto_shapes:
                biggest_shape = max(auto_shapes, key=lambda s: s.width * s.height)
                main_txt = max(text_children, key=lambda t: sum(len(r.text) for r in t.block.runs))
                
                self._embed_text_in_shape(biggest_shape, main_txt)
                
                for txt in text_children:
                    if txt is not main_txt:
                        tb = self._render_text_freeform(txt)
                        if tb:
                            ppt_shapes.append(tb)
                return ppt_shapes
        
        for txt in text_children:
            tb = self._render_text_freeform(txt)
            if tb:
                ppt_shapes.append(tb)
        return ppt_shapes

    def _render_group(self, group: IRGroup):
        ppt_shapes = []
        
        text_children = [c for c in group.children if isinstance(c, IRText)]
        other_children = [c for c in group.children if not isinstance(c, IRText)]

        for child in other_children:
            shape = self._dispatch_shape(child)
            if shape:
                ppt_shapes.append(shape)

        ppt_shapes = self._resolve_group_texts(ppt_shapes, text_children)
                
        ppt_shapes = [s for s in ppt_shapes if s is not None]
        if len(ppt_shapes) > 1:
            try:
                g = self.slide.shapes.add_group_shape(ppt_shapes)
                self._disable_shadow(g)
                result = g
            except Exception as e:
                logger.warning("Could not group shapes: %s", e)
                return None
        elif len(ppt_shapes) == 1:
            result = ppt_shapes[0]
        else:
            return None

        if result and getattr(group, 'shape_id', None):
            self._registry[group.shape_id] = result
            self._connectable_shapes.append(result)

        return result

    def _render_infobox(self, box: IRInfoBox):
        ppt_shapes = []
        
        rect_shape = self._render_rectangle(box.rectangle)
        ppt_shapes.append(rect_shape)

        text_children = []
        other_children = []
        
        for c in box.children:
            if isinstance(c, IRText):
                text_children.append(c)
            else:
                other_children.append(c)

        for child in other_children:
            shape = self._dispatch_shape(child)
            if shape:
                ppt_shapes.append(shape)

        ppt_shapes = self._resolve_group_texts(ppt_shapes, text_children)

        ppt_shapes = [s for s in ppt_shapes if s is not None]
        if len(ppt_shapes) > 1:
            try:
                g = self.slide.shapes.add_group_shape(ppt_shapes)
                self._disable_shadow(g)
                return g
            except Exception as e:
                logger.warning("Could not group shapes: %s", e)
                return None
        elif len(ppt_shapes) == 1:
            return ppt_shapes[0]
        return None

    def _render_rectangle(self, rect: IRRectangle):
        g = rect.geometry
        shape_type = (MSO_SHAPE.ROUNDED_RECTANGLE
                      if rect.corner_radius_emu > 0
                      else MSO_SHAPE.RECTANGLE)
        shape = self.slide.shapes.add_shape(
            shape_type, g.x, g.y, g.width, g.height
        )
        
        if rect.corner_radius_emu > 0 and len(shape.adjustments) > 0:
            min_dim = min(g.width, g.height)
            if min_dim > 0:
                adj_val = min(0.5, rect.corner_radius_emu / min_dim)
                shape.adjustments[0] = adj_val

        if rect.fill:
            if isinstance(rect.fill, GradientFill):
                self._apply_gradient_fill(shape, rect.fill)
            else:
                shape.fill.solid()
                shape.fill.fore_color.rgb = rect.fill.to_rgb()
                self._apply_fill_opacity(shape, rect.fill.opacity)
        else:
            shape.fill.background()

        if rect.stroke:
            shape.line.color.rgb = rect.stroke.to_rgb()
            shape.line.width = max(rect.stroke_width_emu, 1)
            self._apply_line_opacity(shape, rect.stroke.opacity)
        else:
            shape.line.fill.background()

        if rect.dashed:
            self._apply_dash(shape)

        self._disable_shadow(shape)

        shape.text_frame.text = ""
        shape.text_frame.word_wrap = True

        if rect.shape_id:
            self._registry[rect.shape_id] = shape
        self._connectable_shapes.append(shape)
        return shape

    def _render_ellipse(self, el: IREllipse):
        g = el.geometry
        shape = self.slide.shapes.add_shape(
            MSO_SHAPE.OVAL, g.x, g.y, g.width, g.height
        )
        
        if el.fill:
            if isinstance(el.fill, GradientFill):
                self._apply_gradient_fill(shape, el.fill)
            else:
                shape.fill.solid()
                shape.fill.fore_color.rgb = el.fill.to_rgb()
                self._apply_fill_opacity(shape, el.fill.opacity)
        else:
            shape.fill.background()

        if el.stroke:
            shape.line.color.rgb = el.stroke.to_rgb()
            shape.line.width = max(el.stroke_width_emu, 1)
            self._apply_line_opacity(shape, el.stroke.opacity)
        else:
            shape.line.fill.background()

        if el.dashed:
            self._apply_dash(shape)

        self._disable_shadow(shape)
        shape.text_frame.text = ""
        shape.text_frame.word_wrap = True

        if el.shape_id:
            self._registry[el.shape_id] = shape
        self._connectable_shapes.append(shape)
        return shape

    def _render_polygon(self, poly: IRPolygon):
        if not poly.waypoints:
            return None
            
        start = poly.waypoints[0]
        fb = self.slide.shapes.build_freeform(start.x, start.y)
        
        # Add remaining segments
        if len(poly.waypoints) > 1:
            points = [(pt.x, pt.y) for pt in poly.waypoints[1:]]
            fb.add_line_segments(points, close=poly.is_closed)
            
        shape = fb.convert_to_shape()

        if poly.fill and poly.is_closed:
            if isinstance(poly.fill, GradientFill):
                self._apply_gradient_fill(shape, poly.fill)
            else:
                shape.fill.solid()
                shape.fill.fore_color.rgb = poly.fill.to_rgb()
        else:
            shape.fill.background()

        if poly.stroke:
            shape.line.color.rgb = poly.stroke.to_rgb()
            shape.line.width = max(poly.stroke_width_emu, 1)
        else:
            shape.line.fill.background()

        if poly.dashed:
            self._apply_dash(shape)

        self._disable_shadow(shape)
        
        if poly.shape_id:
            self._registry[poly.shape_id] = shape
        self._connectable_shapes.append(shape)
        return shape

    def _render_icon(self, icon: IRIcon):
        g = icon.geometry
        try:
            stream = io.BytesIO(icon.svg_bytes)
            pic = self.slide.shapes.add_picture(
                stream, g.x, g.y, width=g.width, height=g.height
            )
            self._disable_shadow(pic)
            return pic
        except Exception:
            placeholder = self.slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, g.x, g.y, g.width, g.height
            )
            placeholder.fill.background()
            placeholder.line.color.rgb = RGBColor(0x00, 0x7B, 0xFF)
            self._disable_shadow(placeholder)
            return placeholder

    def _render_image(self, img: IRImage):
        g = img.geometry
        try:
            from PIL import Image
            pil_img = Image.open(io.BytesIO(img.image_bytes))
            orig_w, orig_h = pil_img.size

            # Calculate aspect-ratio-preserving dimensions
            aspect = orig_w / orig_h
            target_w = g.width
            target_h = g.height
            target_aspect = target_w / target_h

            if aspect > target_aspect:
                # Image is wider than target - fit to width
                new_w = target_w
                new_h = int(target_w / aspect)
            else:
                # Image is taller than target - fit to height
                new_h = target_h
                new_w = int(target_h * aspect)

            # Center within the target area
            offset_x = (target_w - new_w) // 2
            offset_y = (target_h - new_h) // 2

            stream = io.BytesIO(img.image_bytes)
            pic = self.slide.shapes.add_picture(
                stream, g.x + offset_x, g.y + offset_y, width=new_w, height=new_h
            )
            self._disable_shadow(pic)
            return pic
        except Exception as e:
            logger.warning("Could not render image %s: %s", img.href, e)
            placeholder = self.slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, g.x, g.y, g.width, g.height
            )
            placeholder.fill.background()
            placeholder.line.color.rgb = RGBColor(0xCC, 0x00, 0x00)
            self._disable_shadow(placeholder)
            return placeholder

    def _render_text_freeform(self, txt: IRText):
        block = txt.block
        
        lines = []
        current_line = []
        for run in block.runs:
            if run.line_break_before and current_line:
                lines.append(current_line)
                current_line = []
            current_line.append(run)
        if current_line:
            lines.append(current_line)

        max_width_pt = 0
        total_height_pt = 0
        for line in lines:
            # Add extra space for text width calculation (increased multiplier + padding)
            line_w = sum(len(r.text) * r.font.size_pt * 0.60 for r in line)
            line_h = max((r.font.size_pt for r in line), default=12) * 1.2
            max_width_pt = max(max_width_pt, line_w)
            total_height_pt += line_h

        est_width = Emu(int((max_width_pt + 20) * 12700))
        est_height = Emu(int((total_height_pt + 10) * 12700))

        if txt.container_geometry:
            approx_width = min(est_width, txt.container_geometry.width)
            approx_height = est_height
        else:
            approx_width = min(est_width, Emu(int(self.current_ir_slide.width_emu * 0.9)))
            approx_height = est_height

        if block.alignment == PP_ALIGN.CENTER:
            x = block.anchor_x - Emu(int(approx_width / 2))
        elif block.alignment == PP_ALIGN.RIGHT:
            x = block.anchor_x - approx_width
        else:
            x = block.anchor_x

        y = block.anchor_y - Emu(int(approx_height * 0.5))

        tb = self.slide.shapes.add_textbox(x, y, approx_width, approx_height)
        tb.fill.background()
        tb.line.fill.background()
        self._disable_shadow(tb)
        tf = tb.text_frame
        tf.word_wrap = True
        self._populate_textframe(tf, block)
        return tb

    def _populate_textframe(self, tf, block: TextBlock) -> None:
        para = tf.paragraphs[0]
        para.alignment = block.alignment

        first = True
        for run in block.runs:
            if not first and run.line_break_before:
                para = tf.add_paragraph()
                para.alignment = block.alignment
            first = False

            r = para.add_run()
            r.text = run.text
            f = r.font
            f.name = run.font.family
            f.size = Pt(max(run.font.size_pt, 6))
            f.bold = run.font.bold
            f.italic = run.font.italic
            f.color.rgb = run.font.color.to_rgb()

    def _render_decorative_line(self, line: IRLine):
        conn = self.slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            line.start.x, line.start.y,
            line.end.x, line.end.y,
        )
        conn.line.color.rgb = line.stroke.to_rgb()
        conn.line.width = max(line.stroke_width_emu, Emu(9525))
        self._disable_shadow(conn)
        if line.dashed:
            self._apply_dash(conn)
        return conn

    def _get_closest_connection_site(self, shape, pt: Point) -> int:
        hc = shape.left + shape.width / 2
        vc = shape.top + shape.height / 2
        sites = [
            (0, hc, shape.top),
            (1, shape.left, vc),
            (2, hc, shape.top + shape.height),
            (3, shape.left + shape.width, vc),
        ]
        best = min(sites, key=lambda s: (s[1] - pt.x)**2 + (s[2] - pt.y)**2)
        return best[0]

    def _find_closest_shape(self, pt: Point):
        best_shape = None
        min_dist = float('inf')
        for shape in self._connectable_shapes:
            dx = max(float(shape.left - pt.x), 0.0, float(pt.x - (shape.left + shape.width)))
            dy = max(float(shape.top - pt.y), 0.0, float(pt.y - (shape.top + shape.height)))
            dist = dx**2 + dy**2
            if dist < min_dist:
                min_dist = dist
                best_shape = shape
        
        # Optionally, restrict distance if needed, but since it's an explicit connector we want the closest.
        return best_shape

    def _render_connector(self, c: IRConnector) -> None:
        if len(c.waypoints) < 2:
            return

        is_elbow = c.connector_type == ConnectorType.ELBOW
        conn_type = MSO_CONNECTOR.ELBOW if is_elbow else MSO_CONNECTOR.STRAIGHT

        start_shape = None
        end_shape = None

        if c.start_ref and c.start_ref in self._registry:
            start_shape = self._registry[c.start_ref]
        else:
            start_shape = self._find_closest_shape(c.waypoints[0])

        if c.end_ref and c.end_ref in self._registry:
            end_shape = self._registry[c.end_ref]
        else:
            end_shape = self._find_closest_shape(c.waypoints[-1])

        if start_shape and end_shape:
            start_pt = c.waypoints[0]
            end_pt = c.waypoints[-1]
            
            start_x, start_y = start_pt.x, start_pt.y
            end_x, end_y = end_pt.x, end_pt.y+1
            
            if is_elbow:
                if start_x == end_x:
                    end_x += Emu(10000)
                if start_y == end_y:
                    end_y += Emu(10000)
            
            conn = self.slide.shapes.add_connector(
                conn_type, start_x, start_y, end_x, end_y
            )
            conn.line.color.rgb = c.stroke.to_rgb()
            conn.line.width = max(c.stroke_width_emu, Emu(9525))
            self._disable_shadow(conn)
            if c.dashed:
                self._apply_dash(conn)
            self._apply_arrows(conn, c.start_arrow, c.end_arrow)
            
            conn.begin_connect(shape=start_shape, cxn_pt_idx=self._get_closest_connection_site(start_shape, start_pt))
            conn.end_connect(shape=end_shape, cxn_pt_idx=self._get_closest_connection_site(end_shape, end_pt))
            
        elif c.connector_type == ConnectorType.STRAIGHT or len(c.waypoints) == 2:
            self._draw_segment(
                c.waypoints[0], c.waypoints[-1],
                c.stroke, c.stroke_width_emu, c.dashed,
                start_arrow=c.start_arrow, end_arrow=c.end_arrow,
            )
        else:
            segments = list(zip(c.waypoints[:-1], c.waypoints[1:]))
            last_idx = len(segments) - 1
            for i, (a, b) in enumerate(segments):
                self._draw_segment(
                    a, b, c.stroke, c.stroke_width_emu, c.dashed,
                    start_arrow=c.start_arrow if i == 0 else ArrowType.NONE,
                    end_arrow=c.end_arrow if i == last_idx else ArrowType.NONE,
                )

    def _draw_segment(self, a: Point, b: Point,
                      color: Color, width_emu: int, dashed: bool,
                      start_arrow: ArrowType, end_arrow: ArrowType) -> None:
        conn = self.slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, a.x, a.y, b.x, b.y
        )
        conn.line.color.rgb = color.to_rgb()
        conn.line.width = max(width_emu, Emu(9525))
        self._disable_shadow(conn)
        if dashed:
            self._apply_dash(conn)
        self._apply_arrows(conn, start_arrow, end_arrow)

    def _apply_dash(self, shape) -> None:
        ln = shape.line._get_or_add_ln()
        for child in ln.findall(qn("a:prstDash")):
            ln.remove(child)
        prst = lxml_etree.SubElement(ln, qn("a:prstDash"))
        prst.set("val", "dash")

    def _apply_arrows(self, shape, start: ArrowType, end: ArrowType) -> None:
        ln = shape.line._get_or_add_ln()
        if start == ArrowType.TRIANGLE:
            for c in ln.findall(qn("a:headEnd")):
                ln.remove(c)
            he = lxml_etree.SubElement(ln, qn("a:headEnd"))
            he.set("type", "triangle")
            he.set("w", "med")
            he.set("len", "med")
        if end == ArrowType.TRIANGLE:
            for c in ln.findall(qn("a:tailEnd")):
                ln.remove(c)
            te = lxml_etree.SubElement(ln, qn("a:tailEnd"))
            te.set("type", "triangle")
            te.set("w", "med")
            te.set("len", "med")

    def _disable_shadow(self, shape) -> None:
        try:
            if hasattr(shape, "shadow"):
                shape.shadow.inherit = False
        except Exception:
            pass
        try:
            elem = shape._element
            pr = getattr(elem, "spPr", getattr(elem, "grpSpPr", None))
            if pr is not None:
                for child in pr.findall(qn("a:effectLst")):
                    pr.remove(child)
                lxml_etree.SubElement(pr, qn("a:effectLst"))
        except Exception:
            pass

    def _apply_fill_opacity(self, shape, opacity: float) -> None:
        if opacity >= 1.0:
            return
        try:
            elem = shape._element
            spPr = elem.find(qn("p:spPr"))
            if spPr is None:
                spPr = elem.find(qn("wsp:spPr"))
            if spPr is None:
                return
            solidFill = spPr.find(qn("a:solidFill"))
            if solidFill is None:
                return
            srgbClr = solidFill.find(qn("a:srgbClr"))
            if srgbClr is not None:
                alpha_val = int(opacity * 100000)
                alpha = lxml_etree.SubElement(srgbClr, qn("a:alpha"))
                alpha.set("val", str(alpha_val))
        except Exception:
            pass

    def _apply_line_opacity(self, shape, opacity: float) -> None:
        if opacity >= 1.0:
            return
        try:
            ln = shape.line._get_or_add_ln()
            solidFill = ln.find(qn("a:solidFill"))
            if solidFill is None:
                return
            srgbClr = solidFill.find(qn("a:srgbClr"))
            if srgbClr is not None:
                alpha_val = int(opacity * 100000)
                alpha = lxml_etree.SubElement(srgbClr, qn("a:alpha"))
                alpha.set("val", str(alpha_val))
        except Exception:
            pass

    def _apply_gradient_fill(self, shape, gradient: GradientFill) -> None:
        """Apply a gradient fill to a shape."""
        try:
            shape.fill.gradient()
            # Set gradient angle (PowerPoint uses 60000ths of a degree)
            lin_angle = int(gradient.angle * 60000)
            # Set stops
            stops = shape.fill.gradient_stops
            # Remove existing stops (keep at least one)
            while len(stops) > 1:
                stops[0]._element.getparent().remove(stops[0]._element)
            # Add gradient stops
            gsLst = shape._element.find(qn('p:spPr')).find(qn('a:gradFill')).find(qn('a:gsLst'))
            if gsLst is None:
                gradFill = shape._element.find(qn('p:spPr')).find(qn('a:gradFill'))
                gsLst = lxml_etree.SubElement(gradFill, qn('a:gsLst'))
            # Clear existing stops
            for gs in gsLst.findall(qn('a:gs')):
                gsLst.remove(gs)
            # Add new stops
            for stop in gradient.stops:
                gs = lxml_etree.SubElement(gsLst, qn('a:gs'))
                gs.set('pos', str(int(stop.position * 100000)))
                srgbClr = lxml_etree.SubElement(gs, qn('a:srgbClr'))
                srgbClr.set('val', f'{stop.color.r:02X}{stop.color.g:02X}{stop.color.b:02X}')
            # Set linear angle
            lin = shape._element.find(qn('p:spPr')).find(qn('a:gradFill')).find(qn('a:lin'))
            if lin is None:
                lin = lxml_etree.SubElement(shape._element.find(qn('p:spPr')).find(qn('a:gradFill')), qn('a:lin'))
            lin.set('ang', str(lin_angle))
            lin.set('scaled', '1')
        except Exception as e:
            logger.warning("Could not apply gradient fill: %s", e)
            shape.fill.background()